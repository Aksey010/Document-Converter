"""Shared pytest fixtures and configuration."""

import asyncio
import pytest
import tempfile
import os
import shutil
from pathlib import Path
from unittest.mock import MagicMock, AsyncMock

import fitz
from PIL import Image
from docx import Document

# Add project root to path
import sys
sys.path.insert(0, str(Path(__file__).parent.parent))

from converter import DocumentConverter, ConversionOptions, ConversionResult, OutputFormat, InputFormat
from converter.libraries import LibraryManager
from converter.formats.pdf_handler import PDFHandler
from converter.formats.docx_handler import DOCXHandler
from converter.formats.txt_handler import TXTHandler
from converter.formats.html_handler import HTMLHandler
from converter.formats.image_handler import ImageHandler
from converter.formats.djvu_handler import DJVUHandler
from converter.formats.base import BaseHandler


# ============================================================================
# Session-scoped fixtures
# ============================================================================

@pytest.fixture(scope="session")
def event_loop():
    """Create event loop for async tests."""
    loop = asyncio.new_event_loop()
    yield loop
    loop.close()


@pytest.fixture(scope="session")
def test_dir():
    """Temporary directory for test files."""
    tmpdir = Path(tempfile.mkdtemp(prefix="docconv_test_"))
    yield tmpdir
    shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.fixture(scope="session")
def converter():
    """DocumentConverter instance."""
    return DocumentConverter()


# ============================================================================
# Test file fixtures
# ============================================================================

@pytest.fixture
def sample_text_pdf(test_dir):
    """Create a PDF with extractable text."""
    pdf_path = test_dir / "sample_text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((100, 100), "Test PDF with text content for extraction")
    page.insert_text((100, 150), "Second line of text")
    page.insert_text((100, 200), "Third line with unicode: café, naïve, résumé")
    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def empty_pdf(test_dir):
    """Create an empty PDF (no text)."""
    pdf_path = test_dir / "empty.pdf"
    doc = fitz.open()
    page = doc.new_page()
    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def scanned_pdf(test_dir):
    """Create an image-based PDF (scanned document)."""
    pdf_path = test_dir / "scanned.pdf"
    doc = fitz.open()
    page = doc.new_page()
    img = Image.new('RGB', (200, 200), color='red')
    img_path = test_dir / "temp_img.png"
    img.save(img_path)
    page.insert_image(page.rect, filename=str(img_path))
    doc.save(pdf_path)
    doc.close()
    os.remove(img_path)
    return pdf_path


@pytest.fixture
def sample_docx(test_dir):
    """Create a DOCX with text and table."""
    docx_path = test_dir / "sample.docx"
    doc = Document()
    doc.add_heading("Test Document", level=1)
    doc.add_paragraph("This is a test paragraph with some text.")
    doc.add_paragraph("Another paragraph with different content.")
    
    # Add a table
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(0, 2).text = "Header 3"
    table.cell(1, 0).text = "Row 1 Col 1"
    table.cell(1, 1).text = "Row 1 Col 2"
    table.cell(1, 2).text = "Row 1 Col 3"
    table.cell(2, 0).text = "Row 2 Col 1"
    table.cell(2, 1).text = "Row 2 Col 2"
    table.cell(2, 2).text = "Row 2 Col 3"
    
    doc.save(docx_path)
    return docx_path


@pytest.fixture
def sample_txt(test_dir):
    """Create a text file."""
    txt_path = test_dir / "sample.txt"
    content = "Test TXT content\nLine 2\nLine 3 with unicode: café, naïve, résumé"
    txt_path.write_text(content, encoding='utf-8')
    return txt_path


@pytest.fixture
def sample_html(test_dir):
    """Create an HTML file."""
    html_path = test_dir / "sample.html"
    content = """<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>Test HTML</title>
</head>
<body>
    <h1>Test HTML Document</h1>
    <p>This is a test paragraph.</p>
    <ul>
        <li>Item 1</li>
        <li>Item 2</li>
        <li>Item 3</li>
    </ul>
    <table>
        <tr><th>Header 1</th><th>Header 2</th></tr>
        <tr><td>Cell 1</td><td>Cell 2</td></tr>
    </table>
</body>
</html>"""
    html_path.write_text(content, encoding='utf-8')
    return html_path


@pytest.fixture
def sample_image(test_dir):
    """Create a test image (PNG)."""
    img_path = test_dir / "sample.png"
    img = Image.new('RGB', (200, 200), color='red')
    # Add some pattern
    for x in range(200):
        for y in range(200):
            if (x + y) % 20 < 10:
                img.putpixel((x, y), (0, 255, 0))
    img.save(img_path)
    return img_path


@pytest.fixture
def sample_jpeg(test_dir):
    """Create a test JPEG."""
    img_path = test_dir / "sample.jpg"
    img = Image.new('RGB', (200, 200), color='blue')
    img.save(img_path, 'JPEG', quality=85)
    return img_path


# ============================================================================
# Handler fixtures
# ============================================================================

@pytest.fixture
def pdf_handler():
    """PDF handler instance."""
    return PDFHandler()


@pytest.fixture
def docx_handler():
    """DOCX handler instance."""
    return DOCXHandler()


@pytest.fixture
def txt_handler():
    """TXT handler instance."""
    return TXTHandler()


@pytest.fixture
def html_handler():
    """HTML handler instance."""
    return HTMLHandler()


@pytest.fixture
def image_handler():
    """Image handler instance."""
    return ImageHandler()


@pytest.fixture
def djvu_handler():
    """DJVU handler instance."""
    return DJVUHandler()


@pytest.fixture
def library_manager():
    """Library manager instance."""
    return LibraryManager()


# ============================================================================
# Mock fixtures
# ============================================================================

@pytest.fixture
def mock_fitz(monkeypatch):
    """Mock fitz module."""
    mock = MagicMock()
    mock.open = MagicMock()
    monkeypatch.setattr('fitz.open', mock.open)
    return mock


# ============================================================================
# Helper functions
# ============================================================================

def create_multipart_form_data(file_path: Path, output_format: str, options: dict = None):
    """Create multipart form data for API testing."""
    import uuid
    boundary = f'----WebKitFormBoundary{uuid.uuid4().hex[:16]}'
    
    with open(file_path, 'rb') as f:
        file_data = f.read()
    
    parts = [
        f'--{boundary}',
        f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"',
        f'Content-Type: application/octet-stream',
        '',
    ]
    body = '\r\n'.join(parts).encode() + b'\r\n' + file_data + b'\r\n'
    body += f'--{boundary}\r\n'.encode()
    body += b'Content-Disposition: form-data; name="output_format"\r\n\r\n'
    body += output_format.encode() + b'\r\n'
    body += f'--{boundary}\r\n'.encode()
    body += b'Content-Disposition: form-data; name="options"\r\n\r\n'
    body += json.dumps(options or {}).encode() + b'\r\n'
    body += f'--{boundary}--\r\n'.encode()
    
    return body, boundary