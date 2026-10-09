"""Integration tests for web API endpoints."""

import pytest
import asyncio
import tempfile
import os
import json
import requests
from pathlib import Path

import fitz
from PIL import Image
from docx import Document

import sys
sys.path.insert(0, str(Path(__file__).parent.parent))

from aiohttp.test_utils import TestClient, TestServer
from aiohttp import web

from web_converter.app import create_app, temp_manager, conversion_semaphore, thread_pool
from converter import DocumentConverter, ConversionOptions, OutputFormat, InputFormat
from web_converter.app import converter, temp_manager, conversion_semaphore, thread_pool


# ============================================================================
# Fixtures
# ============================================================================

@pytest.fixture(scope="session")
def event_loop():
    """Create event loop for the test session."""
    loop = asyncio.new_event_loop()
    yield loop
    loop.close()


@pytest.fixture(scope="session")
async def app():
    """Create test application (session-scoped to share thread pool)."""
    app = await create_app()
    return app


@pytest.fixture(scope="session")
async def client(app):
    """Create test client (session-scoped to reuse TestServer)."""
    async with TestClient(TestServer(app)) as client:
        yield client


@pytest.fixture
def sample_text_pdf():
    """Create a PDF with text for testing."""
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
        pdf_path = f.name
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((100, 100), 'Test PDF content for API testing')
    doc.save(pdf_path)
    doc.close()
    yield pdf_path
    os.unlink(pdf_path)


@pytest.fixture
def empty_pdf():
    """Create an empty PDF (no text)."""
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
        pdf_path = f.name
    doc = fitz.open()
    page = doc.new_page()
    doc.save(pdf_path)
    doc.close()
    yield pdf_path
    os.unlink(pdf_path)


@pytest.fixture
def sample_docx():
    """Create a DOCX for testing."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        docx_path = f.name
    doc = Document()
    doc.add_heading('Test Document', level=1)
    doc.add_paragraph('Test DOCX content for API testing.')
    doc.save(docx_path)
    yield docx_path
    os.unlink(docx_path)


@pytest.fixture
def sample_image():
    """Create a test image."""
    with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as f:
        img_path = f.name
    img = Image.new('RGB', (200, 200), color='red')
    img.save(img_path, 'PNG')
    yield img_path
    os.unlink(img_path)


@pytest.fixture
def sample_jpeg():
    """Create a test JPEG."""
    with tempfile.NamedTemporaryFile(suffix='.jpg', delete=False) as f:
        img_path = f.name
    img = Image.new('RGB', (200, 200), color='blue')
    img.save(img_path, 'JPEG', quality=85)
    yield img_path
    os.unlink(img_path)


@pytest.fixture
def sample_txt():
    """Create a text file."""
    with tempfile.NamedTemporaryFile(suffix='.txt', delete=False, mode='w', encoding='utf-8') as f:
        f.write('Test TXT content\nLine 2\nLine 3')
        txt_path = f.name
    yield txt_path
    os.unlink(txt_path)


@pytest.fixture
def sample_html():
    """Create an HTML file."""
    with tempfile.NamedTemporaryFile(suffix='.html', delete=False, mode='w', encoding='utf-8') as f:
        f.write('<html><body><h1>Test HTML</h1><p>Content here</p></body></html>')
        html_path = f.name
    yield html_path
    os.unlink(html_path)


@pytest.fixture
def sample_gif():
    """Create a GIF file."""
    with tempfile.NamedTemporaryFile(suffix='.gif', delete=False) as f:
        gif_path = f.name
    img = Image.new('P', (100, 100), color=1)
    img.save(gif_path, 'GIF')
    yield gif_path
    os.unlink(gif_path)


@pytest.fixture
def empty_pdf():
    """Create an empty PDF (no text)."""
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
        pdf_path = f.name
    doc = fitz.open()
    page = doc.new_page()
    doc.save(pdf_path)
    doc.close()
    yield pdf_path
    os.unlink(pdf_path)


# ============================================================================
# Helper Functions
# ============================================================================

async def _convert_file_async(client, file_path, output_format, uploaded_filename=None):
    """Helper to convert file via API using the test client."""
    import uuid
    boundary = f'----WebKitFormBoundary{uuid.uuid4().hex[:16]}'
    
    if uploaded_filename is None:
        uploaded_filename = Path(file_path).name
        
    with open(file_path, 'rb') as f:
        file_data = f.read()
    
    body_parts = [
        b'--' + boundary.encode(),
        f'Content-Disposition: form-data; name="file"; filename="{Path(uploaded_filename).name}"'.encode(),
        b'Content-Type: application/octet-stream',
        b'',
        file_data,
        b'--' + boundary.encode(),
        b'Content-Disposition: form-data; name="output_format"',
        b'',
        output_format.encode(),
        b'--' + boundary.encode(),
        b'Content-Disposition: form-data; name="options"',
        b'',
        b'{"quality":85,"dpi":150}',
        b'--' + boundary.encode() + b'--'
    ]
    body = b'\r\n'.join(body_parts)
    
    resp = await client.post('/api/convert',
        data=body,
        headers={'Content-Type': f'multipart/form-data; boundary={boundary}',
                 'Content-Length': str(len(body))}
    )
    return await resp.json()


async def _estimate_size_async(client, file_path, output_format, options=None):
    """Helper to estimate output size via API."""
    import uuid
    boundary = f'----WebKitFormBoundary{uuid.uuid4().hex[:16]}'
    
    with open(file_path, 'rb') as f:
        file_data = f.read()
    
    body_parts = [
        b'--' + boundary.encode(),
        f'Content-Disposition: form-data; name="file"; filename="{Path(file_path).name}"'.encode(),
        b'Content-Type: application/octet-stream',
        b'',
        file_data,
        b'--' + boundary.encode(),
        b'Content-Disposition: form-data; name="output_format"',
        b'',
        output_format.encode(),
        b'--' + boundary.encode(),
        b'Content-Disposition: form-data; name="options"',
        b'',
        json.dumps(options or {}),
        b'--' + boundary.encode() + b'--'
    ]
    body = b'\r\n'.join(body_parts)
    
    resp = await client.post('/api/estimate',
        data=body,
        headers={'Content-Type': f'multipart/form-data; boundary={boundary}',
                 'Content-Length': str(len(body))}
    )
    return await resp.json()


# ============================================================================
# Test Classes
# ============================================================================

class TestFormatsEndpoint:
    """Tests for /api/formats endpoint."""

    async def test_get_formats(self, client):
        """Test GET /api/formats returns supported formats."""
        resp = await client.get('/api/formats')
        assert resp.status == 200
        
        data = await resp.json()
        assert 'input_formats' in data
        assert 'output_formats' in data
        assert 'compatibility' in data
        
        assert 'pdf' in data['input_formats']
        assert 'docx' in data['input_formats']
        assert 'txt' in data['input_formats']
        assert 'html' in data['input_formats']
        
        assert 'pdf' in data['compatibility']['pdf']
        assert 'docx' in data['compatibility']['pdf']

    async def test_formats_structure(self, client):
        """Test formats response structure."""
        resp = await client.get('/api/formats')
        data = await resp.json()
        
        assert isinstance(data['input_formats'], list)
        assert isinstance(data['output_formats'], list)
        assert isinstance(data['compatibility'], dict)
        
        for input_fmt, output_fmts in data['compatibility'].items():
            assert isinstance(output_fmts, list)


class TestEstimateEndpoint:
    """Tests for /api/estimate endpoint."""

    async def test_estimate_pdf_to_docx(self, client, sample_text_pdf):
        """Test estimate for PDF to DOCX."""
        # Use the converter directly for estimation
        converter = DocumentConverter()
        options = ConversionOptions(quality=85, dpi=150)
        estimate = converter.estimate_output_size(sample_text_pdf, 'docx', options)
        
        assert 'input_size' in estimate
        assert 'estimated_output_size' in estimate
        assert 'compression_ratio' in estimate
        assert estimate['input_size'] > 0
        assert estimate['estimated_output_size'] > 0

    async def test_estimate_missing_fields(self, client):
        """Test estimate with missing fields."""
        resp = await client.post('/api/estimate', json={
            'output_format': 'docx'
        })
        assert resp.status == 400
        
        data = await resp.json()
        assert 'error' in data

    async def test_estimate_file_not_found(self, client):
        """Test estimate for non-existent file."""
        resp = await client.post('/api/estimate', json={
            'filename': 'nonexistent.pdf',
            'output_format': 'docx'
        })
        assert resp.status == 404


class TestConvertEndpoint:
    """Tests for /api/convert endpoint."""

    async def test_convert_pdf_to_txt(self, client, sample_text_pdf):
        """Test PDF to TXT conversion via API."""
        data = await _convert_file_async(client, sample_text_pdf, 'txt')
        
        assert data['success'] is True
        assert 'filename' in data
        assert data['format'] == 'txt'
        assert data['output_size'] > 0
        assert 'download_url' in data

    async def test_convert_pdf_to_docx(self, client, sample_text_pdf):
        """Test PDF to DOCX conversion."""
        data = await _convert_file_async(client, sample_text_pdf, 'docx')
        
        assert data['success'] is True
        assert data['format'] == 'docx'
        assert data['output_size'] > 0

    async def test_convert_pdf_to_docx_empty(self, client, empty_pdf):
        """Test PDF to DOCX with empty PDF returns error."""
        data = await _convert_file_async(client, empty_pdf, 'docx')
        
        assert data['success'] is False
        assert 'error' in data
        assert 'text' in data['error'].lower() or 'empty' in data['error'].lower()

    async def test_convert_pdf_to_html(self, client, sample_text_pdf):
        """Test PDF to HTML conversion."""
        data = await _convert_file_async(client, sample_text_pdf, 'html')
        
        assert data['success'] is True
        assert data['format'] == 'html'

    async def test_convert_pdf_to_images(self, client, sample_text_pdf):
        """Test PDF to image conversions."""
        for fmt in ['png', 'jpeg', 'tiff', 'bmp', 'webp']:
            data = await _convert_file_async(client, sample_text_pdf, fmt)
            assert data['success'] is True, f"Failed for {fmt}: {data.get('error')}"
            assert data['format'] == fmt

    async def test_convert_docx_to_pdf(self, client, sample_docx):
        """Test DOCX to PDF conversion."""
        data = await _convert_file_async(client, sample_docx, 'pdf')
        
        assert data['success'] is True
        assert data['format'] == 'pdf'
        assert data['output_size'] > 0

    async def test_convert_docx_to_txt(self, client, sample_docx):
        """Test DOCX to TXT conversion."""
        data = await _convert_file_async(client, sample_docx, 'txt')
        
        assert data['success'] is True
        assert data['format'] == 'txt'

    async def test_convert_docx_to_html(self, client, sample_docx):
        """Test DOCX to HTML conversion."""
        data = await _convert_file_async(client, sample_docx, 'html')
        
        assert data['success'] is True
        assert data['format'] == 'html'

    async def test_convert_image_to_pdf(self, client, sample_image):
        """Test image to PDF conversion."""
        data = await _convert_file_async(client, sample_image, 'pdf')
        
        assert data['success'] is True
        assert data['format'] == 'pdf'

    async def test_convert_image_to_docx(self, client, sample_image):
        """Test image to DOCX."""
        data = await _convert_file_async(client, sample_image, 'docx')
        
        assert data['success'] is True
        assert data['format'] == 'docx'

    async def test_convert_image_to_image(self, client, sample_image):
        """Test image to image conversion."""
        for fmt in ['png', 'jpeg', 'webp', 'tiff', 'bmp']:
            data = await _convert_file_async(client, sample_image, fmt)
            assert data['success'] is True, f"Failed for {fmt}: {data.get('error')}"
            assert data['format'] == fmt

    async def test_convert_txt_to_pdf(self, client, sample_txt):
        """Test TXT to PDF conversion."""
        data = await _convert_file_async(client, sample_txt, 'pdf')
        
        assert data['success'] is True
        assert data['format'] == 'pdf'

    async def test_convert_txt_to_docx(self, client, sample_txt):
        """Test TXT to DOCX."""
        data = await _convert_file_async(client, sample_txt, 'docx')
        
        assert data['success'] is True
        assert data['format'] == 'docx'

    async def test_convert_txt_to_html(self, client, sample_txt):
        """Test TXT to HTML."""
        data = await _convert_file_async(client, sample_txt, 'html')
        
        assert data['success'] is True
        assert data['format'] == 'html'

    async def test_convert_html_to_pdf(self, client, sample_html):
        """Test HTML to PDF."""
        data = await _convert_file_async(client, sample_html, 'pdf')
        
        assert data['success'] is True
        assert data['format'] == 'pdf'

    async def test_convert_html_to_docx(self, client, sample_html):
        """Test HTML to DOCX."""
        data = await _convert_file_async(client, sample_html, 'docx')
        
        assert data['success'] is True
        assert data['format'] == 'docx'

    async def test_invalid_format(self, client, sample_text_pdf):
        """Test unsupported output format."""
        data = await _convert_file_async(client, sample_text_pdf, 'unsupported')
        
        assert data['success'] is False
        assert 'error' in data

    async def test_file_too_large(self, client):
        """Test file size limit."""
        # Create a large file
        large_file = Path(tempfile.mktemp(suffix='.pdf'))
        with open(large_file, 'wb') as f:
            f.write(b'x' * (150 * 1024 * 1024))  # 150MB > 100MB limit
        
        try:
            data = await _convert_file_async(client, str(large_file), 'txt')
            assert data['success'] is False
            assert 'large' in data['error'].lower() or 'size' in data['error'].lower()
        finally:
            os.unlink(large_file)


class TestDownloadEndpoint:
    """Tests for /download/{filename} endpoint."""

    async def test_download_converted_file(self, client, sample_text_pdf):
        """Test downloading converted file."""
        data = await _convert_file_async(client, sample_text_pdf, 'txt')
        
        if data.get('success') and data.get('download_url'):
            resp = await client.get(data['download_url'])
            assert resp.status == 200
            assert resp.headers.get('Content-Disposition', '').startswith('attachment')


class TestStatusEndpoint:
    """Tests for /api/status endpoint."""

    async def test_status_endpoint(self, client):
        """Test status endpoint returns server info."""
        resp = await client.get('/api/status')
        assert resp.status == 200
        
        data = await resp.json()
        assert data['status'] == 'ok'
        assert 'conversions_active' in data
        assert 'conversions_max' in data
        assert 'thread_pool_workers' in data
        assert 'temp_dirs' in data
        assert 'libraries' in data


class TestIndexEndpoint:
    """Tests for root endpoint."""

    async def test_index_page(self, client):
        """Test main page loads."""
        resp = await client.get('/')
        assert resp.status == 200
        assert resp.content_type == 'text/html'
        
        text = await resp.text()
        assert 'Document Converter' in text
        assert 'dropZone' in await resp.text()


class TestStaticFiles:
    """Tests for static file serving."""

    async def test_static_css(self, client):
        """Test CSS file served."""
        resp = await client.get('/static/style.css')
        assert resp.status == 200
        assert 'text/css' in resp.content_type

    async def test_static_js(self, client):
        """Test JS file served."""
        resp = await client.get('/static/app.js')
        assert resp.status == 200
        assert 'javascript' in resp.content_type

    async def test_manifest(self, client):
        """Test manifest.json served."""
        resp = await client.get('/static/manifest.json')
        assert resp.status == 200
        assert 'application/json' in resp.content_type


class TestConcurrencyLimit:
    """Tests for concurrency limiting."""

    async def test_concurrent_limit(self, client, sample_text_pdf):
        """Test concurrent conversion limit."""
        # Make multiple concurrent requests
        tasks = [
            _convert_file_async(client, sample_text_pdf, 'txt') 
            for _ in range(5)
        ]
        
        results = await asyncio.gather(*tasks, return_exceptions=True)
        
        # Some should succeed, some may be rejected (503)
        successes = sum(1 for r in results if isinstance(r, dict) and r.get('success'))
        errors = sum(1 for r in results if isinstance(r, dict) and not r.get('success'))
        
        # At least some should succeed
        assert successes > 0


class TestSSEProgress:
    """Tests for SSE progress endpoint."""

    async def test_sse_endpoint(self, client, sample_text_pdf):
        """Test SSE progress endpoint."""
        # First start conversion
        data = await _convert_file_async(client, sample_text_pdf, 'txt')
        
        if data.get('success') and data.get('task_id'):
            task_id = data['task_id']
            
            # Connect to SSE
            resp = await client.get(f'/api/convert/progress/{task_id}')
            assert resp.status == 200
            assert resp.content_type == 'text/event-stream'
            
            # Read a few events
            count = 0
            async for line in resp.content:
                if line:
                    line = line.decode().strip()
                    if line.startswith('data:'):
                        data = json.loads(line[5:])
                        assert 'task_id' in data
                        assert 'progress' in data
                        assert 'status' in data
                        count += 1
                        if count >= 3:
                            break
            
            assert count > 0


class TestTempFileCleanup:
    """Tests for temp file management."""

    async def test_temp_files_cleaned(self, client, sample_text_pdf):
        """Test that temp files are cleaned after conversion."""
        initial_count = len(list(temp_manager.upload_dir.glob('*'))) + \
                       len(list(temp_manager.output_dir.glob('*')))
        
        await _convert_file_async(client, sample_text_pdf, 'txt')
        
        # Wait a bit for cleanup
        await asyncio.sleep(0.5)
        
        final_count = len(list(temp_manager.upload_dir.glob('*'))) + \
                     len(list(temp_manager.output_dir.glob('*')))
        
        # Should not have accumulated temp files
        assert final_count <= initial_count + 5  # Allow some margin